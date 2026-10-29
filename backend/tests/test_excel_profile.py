from pathlib import Path
import openpyxl
import pandas as pd


def test_source_file_and_sheet_count():
    """Verify source Excel file exists and contains all 14 expected sheets."""
    base_dir = Path(__file__).resolve().parent.parent.parent
    file_path = base_dir / "data" / "source" / "KA rentals DB.xlsx"
    assert file_path.exists(), "Source workbook missing in data/source"

    wb = openpyxl.load_workbook(file_path, read_only=True)
    expected_sheets = [
        "Customer details",
        "Sheet15",
        "paying needed",
        "pay day",
        "Current Contracts",
        "Milegae",
        "Sheet12",
        "Outstanding PCN",
        "kashif ",
        "Mohammed PCN",
        "Mohammed COW PCN",
        "Adangoh ",
        "Ashley ",
        "Edwards PCN",
    ]
    assert len(wb.sheetnames) == 14
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Expected sheet '{s}' missing from workbook"


def test_core_sheets_dimensions():
    """Verify non-empty row counts and column structure for main operational sheets."""
    base_dir = Path(__file__).resolve().parent.parent.parent
    file_path = base_dir / "data" / "source" / "KA rentals DB.xlsx"

    # Customer details
    cust_df = pd.read_excel(file_path, sheet_name="Customer details")
    assert len(cust_df) == 52
    assert "Name" in cust_df.columns
    assert "Email" in cust_df.columns

    # Current Contracts
    cc_df = pd.read_excel(file_path, sheet_name="Current Contracts")
    assert len(cc_df) == 13
    assert "Reg number" in cc_df.columns
    assert "Status contract" in cc_df.columns

    # Outstanding PCN
    pcn_df = pd.read_excel(file_path, sheet_name="Outstanding PCN")
    assert len(pcn_df) == 45
    assert "PCN number" in pcn_df.columns
    assert "Status" in pcn_df.columns


def test_vehicle_registrations_cross_sheet_consistency():
    """Verify that normalized vehicle registrations match across contracts and pay day."""
    base_dir = Path(__file__).resolve().parent.parent.parent
    file_path = base_dir / "data" / "source" / "KA rentals DB.xlsx"

    cc_df = pd.read_excel(file_path, sheet_name="Current Contracts")
    pay_df = pd.read_excel(file_path, sheet_name="pay day")
    s12_df = pd.read_excel(file_path, sheet_name="Sheet12")

    cc_regs = {str(r).replace(" ", "").upper() for r in cc_df["Reg number"].dropna()}
    pay_regs = {str(r).replace(" ", "").upper() for r in pay_df["Registration"].dropna()}
    s12_regs = {str(r).replace(" ", "").upper() for r in s12_df["Car Reg"].dropna()}

    assert len(cc_regs) == 13
    assert cc_regs == pay_regs, "Mismatch between Current Contracts and pay day vehicles"
    assert cc_regs == s12_regs, "Mismatch between Current Contracts and Sheet12 vehicles"


def test_profile_documentation_privacy():
    """Verify docs/excel-data-profile.md exists and does NOT contain private emails."""
    base_dir = Path(__file__).resolve().parent.parent.parent
    doc_path = base_dir / "docs" / "excel-data-profile.md"
    assert doc_path.exists(), "Profile documentation missing"

    doc_text = doc_path.read_text(encoding="utf-8")

    # Load actual emails from Customer details and verify none appear in the docs
    file_path = base_dir / "data" / "source" / "KA rentals DB.xlsx"
    cust_df = pd.read_excel(file_path, sheet_name="Customer details")
    emails = cust_df["Email"].dropna().astype(str).tolist()

    for email in emails:
        if "@" in email and len(email.strip()) > 5:
            assert email.strip().lower() not in doc_text.lower(), (
                f"Privacy violation: email found in documentation: {email}"
            )
